#!/usr/bin/env python3
"""Печатная форма акта: xlsx (лист 1) -> HTML -> PDF/PNG через Chromium. Формулы шаблона просчитываются здесь."""
import sys, re, html, json, subprocess, os, datetime as dt
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter as L

src = sys.argv[1]; out_dir = sys.argv[2] if len(sys.argv) > 2 else os.path.dirname(src)
wb = load_workbook(src); ws = wb.worksheets[0]
MAXR, MAXC = 83, 21  # B..U

def v(c):  # raw value
    return ws[c].value

def n(c):
    x = val(c); return float(x) if isinstance(x, (int, float)) else 0.0

_cache = {}
def val(c):
    """значение с вычислением известных формул шаблона"""
    if c in _cache: return _cache[c]
    x = v(c); r = None
    if isinstance(x, str) and x.startswith('='):
        f = x[1:]
        m = re.fullmatch(r'H(\d+)-G(\d+)', f)
        if m: r = n('H'+m.group(1)) - n('G'+m.group(2))
        m = re.fullmatch(r'J(\d+)\*I(\d+)', f)
        if m and r is None: r = n('J'+m.group(1)) * n('I'+m.group(2))
        m = re.fullmatch(r'SUM\(([A-Z]+)(\d+):([A-Z]+)(\d+)\)', f)
        if m and r is None: r = sum(n(f'{m.group(1)}{i}') for i in range(int(m.group(2)), int(m.group(4))+1))
        if f == 'F19': r = val('F19')
        if f == 'K66+K40': r = n('K66') + n('K40')
        if f == 'K67-K66': r = n('K67') - n('K66')
        if f.startswith('IF(K68<0'): r = 'евро в пользу Ритет' if n('K68') < 0 else 'евро в пользу Эмас'
        if f == 'F16': r = val('F16')
        if f == 'U79/U78': r = n('U79')/n('U78') if n('U78') else 0
        m = re.fullmatch(r'([A-Z])71\+([A-Z])71\+([A-Z])71\+([A-Z])71\+([A-Z])71\+([A-Z])71', f)
        if m: r = sum(n(g+'71') for g in m.groups())
        m = re.fullmatch(r'([A-Z])72\+([A-Z])72\+([A-Z])72\+([A-Z])72\+([A-Z])72\+([A-Z])72', f)
        if m: r = sum(n(g+'72') for g in m.groups())
        m = re.fullmatch(r'\(([A-Z])71\*U79\+([A-Z])72\)\*100/\(D73\*U79\)', f)
        if m: d = n('D73')*n('U79'); r = (n(m.group(1)+'71')*n('U79') + n(m.group(2)+'72'))*100/d if d else 0
        if r is None: r = ''
    else:
        r = x
    _cache[c] = r; return r

def fmt(c):
    x = val(c); cell = ws[c]; nf = cell.number_format or 'General'
    if x is None or x == '': return ''
    if isinstance(x, (dt.datetime, dt.date)): return x.strftime('%d.%m.%Y')
    if isinstance(x, (int, float)):
        if 'mm-dd' in nf or 'dd' in nf: return str(x)
        if isinstance(x, float) and x == 0 and c[0] in 'IK' and 28 <= int(c[1:]) <= 39: return ''
        if nf.startswith('0.0000'): return f'{x:.4f}'.replace('.', ',')
        if nf in ('0.00',) or '0.00' in nf: s = f'{x:,.2f}'
        elif nf.startswith('#,##0') and '.00' not in nf: s = f'{x:,.0f}'
        elif isinstance(x, float) and not x.is_integer(): s = f'{x:,.2f}'
        else: s = f'{x:,.0f}' if float(x).is_integer() else f'{x:,.2f}'
        return s.replace(',', ' ').replace('.', ',')
    return html.escape(str(x)).replace('\n', '<br>')

# сетка: объединения
merged = {}; hidden = set()
for rg in ws.merged_cells.ranges:
    merged[(rg.min_row, rg.min_col)] = (rg.max_row-rg.min_row+1, rg.max_col-rg.min_col+1)
    for r in range(rg.min_row, rg.max_row+1):
        for c in range(rg.min_col, rg.max_col+1):
            if (r, c) != (rg.min_row, rg.min_col): hidden.add((r, c))

def px_w(col):
    w = ws.column_dimensions[L(col)].width or ws.sheet_format.defaultColWidth or 9.14
    return round(w * 7 + 5)
def px_h(row):
    h = ws.row_dimensions[row].height or ws.sheet_format.defaultRowHeight or 15
    return round(h * 96/72)

def border_css(cell, rs, cs, r, c):
    # границы берём с внешних ячеек объединённой области
    b = cell.border; css = []
    def st(side): return 'solid' if side and side.style else None
    top = st(b.top); left = st(b.left)
    br = ws.cell(r+rs-1, c+cs-1).border; bottom = st(br.bottom); right = st(br.right)
    # двойные/толстые не различаем
    for name, s in (('top', top), ('left', left), ('bottom', bottom), ('right', right)):
        if s: css.append(f'border-{name}:1px solid #000')
    return ';'.join(css)

cols_html = ''.join(f'<col style="width:{px_w(c)}px">' for c in range(2, MAXC+1))
rows = []
for r in range(1, MAXR+1):
    if ws.row_dimensions[r].hidden: continue
    tds = []
    for c in range(2, MAXC+1):
        if (r, c) in hidden: continue
        cell = ws.cell(r, c); rs, cs = merged.get((r, c), (1, 1)); coord = f'{L(c)}{r}'
        rs = sum(1 for rr in range(r, r+rs) if not ws.row_dimensions[rr].hidden) or 1   # rowspan по видимым строкам
        f = cell.font; a = cell.alignment
        style = [border_css(cell, rs, cs, r, c)]
        style.append(f'font-size:{(f.sz or 12)*0.78:.1f}pt')
        if f.bold: style.append('font-weight:700')
        if f.italic: style.append('font-style:italic')
        ha = a.horizontal or ('right' if isinstance(val(coord), (int, float)) else 'left')
        if ha == 'general': ha = 'left'
        style.append(f'text-align:{ha}')
        style.append(f'vertical-align:{ {"center":"middle","top":"top","bottom":"bottom"}.get(a.vertical or "bottom","bottom") }')
        if not a.wrap_text: style.append('white-space:nowrap')
        if a.indent: style.append(f'padding-left:{4+a.indent*8}px')
        tds.append(f'<td rowspan="{rs}" colspan="{cs}" style="{";".join(s for s in style if s)}">{fmt(coord)}</td>')
    rows.append(f'<tr style="height:{px_h(r)}px">{"".join(tds)}</tr>')

title = ws.title
doc = f'''<!doctype html><html><head><meta charset="utf-8"><title>{html.escape(title)}</title>
<style>
@page {{ size: A4 landscape; margin: 8mm; }}
body {{ margin:0; font-family:"Times New Roman", Times, serif; color:#000; background:#fff }}
table {{ border-collapse:collapse; table-layout:fixed; }}
td {{ padding:0 3px; overflow:hidden; line-height:1.15 }}
</style></head><body><table>{cols_html}{''.join(rows)}</table></body></html>'''
os.makedirs(out_dir, exist_ok=True)
base = os.path.join(out_dir, re.sub(r'\.xlsx$', '', os.path.basename(src)))
open(base + '.html', 'w', encoding='utf-8').write(doc)
js = f'''
const {{ chromium }} = require('/opt/node22/lib/node_modules/playwright/index.js');
(async () => {{
  const b = await chromium.launch(); const p = await b.newPage({{ viewport: {{ width: 1600, height: 1100 }}, deviceScaleFactor: 2 }});
  await p.goto('file://' + {json.dumps(base + '.html')}); await p.emulateMedia({{ media: 'print' }});
  const bb = await p.evaluate(() => {{ const r = document.querySelector('table').getBoundingClientRect(); return {{ w: r.width, h: r.height }}; }});
  const w = bb.w;
  // A4 landscape: 297×210 мм ≈ 1123×794 px @96dpi, минус поля 8 мм с каждой стороны — вписываем по обеим осям (как fitToPage в Excel)
  const scale = Math.max(0.1, Math.min(1, (1123 - 61) / bb.w, (794 - 61) / bb.h));
  await p.pdf({{ path: {json.dumps(base + '.pdf')}, format: 'A4', landscape: true, printBackground: true, scale, margin: {{ top: '8mm', bottom: '8mm', left: '8mm', right: '8mm' }} }});
  await p.setViewportSize({{ width: Math.ceil(w) + 40, height: 900 }});
  await p.screenshot({{ path: {json.dumps(base + '.png')}, fullPage: true }});
  await b.close(); console.log('scale', scale.toFixed(3));
}})();
'''
open(base + '.render.cjs', 'w').write(js)
print(subprocess.run(['node', base + '.render.cjs'], capture_output=True, text=True).stdout.strip())
os.remove(base + '.render.cjs')
print('written:', base + '.pdf', '|', base + '.png')
